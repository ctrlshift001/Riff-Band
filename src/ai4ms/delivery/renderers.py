from __future__ import annotations

import json
import math
import os
import re
from html import escape
from pathlib import Path
from typing import Any


STATUS_COLORS = {
    "approved": "#16856b",
    "needs_review": "#d48a16",
    "in_progress": "#2878b5",
    "blocked": "#b74343",
    "not_started": "#aeb6bd",
}


def build_visual_assets(project: dict[str, Any], export_dir: Path) -> dict[str, Any]:
    charts_dir = export_dir / "charts"
    diagrams_dir = export_dir / "diagrams"
    charts_dir.mkdir(parents=True, exist_ok=True)
    diagrams_dir.mkdir(parents=True, exist_ok=True)

    stages = [
        {
            "code": str(stage.get("code") or ""),
            "title": str(stage.get("title") or ""),
            "status": str(stage.get("status") or "not_started"),
            "revision": int(stage.get("revision") or 0),
        }
        for stage in project.get("stages", [])
        if isinstance(stage, dict)
    ]
    stage_spec = {
        "chart_id": "stage-revisions",
        "kind": "progress",
        "title": "研究阶段与版本状态",
        "items": stages,
    }
    _write_json(charts_dir / "stage-revisions.json", stage_spec)
    (charts_dir / "stage-revisions.svg").write_text(
        _stage_chart_svg(stage_spec), encoding="utf-8"
    )
    _write_progress_png(charts_dir / "stage-revisions.png", stages)

    chart_records = [
        {
            "chart_id": "stage-revisions",
            "kind": "progress",
            "title": stage_spec["title"],
            "json_path": "charts/stage-revisions.json",
            "svg_path": "charts/stage-revisions.svg",
            "png_path": "charts/stage-revisions.png",
        }
    ]

    estimates = _collect_estimates(project)
    if estimates:
        estimate_spec = {
            "chart_id": "estimate-results",
            "kind": "forest",
            "title": "结构化估计结果与置信区间",
            "items": estimates,
        }
        _write_json(charts_dir / "estimate-results.json", estimate_spec)
        (charts_dir / "estimate-results.svg").write_text(
            _estimate_chart_svg(estimate_spec), encoding="utf-8"
        )
        _write_estimate_png(charts_dir / "estimate-results.png", estimates)
        chart_records.append(
            {
                "chart_id": "estimate-results",
                "kind": "forest",
                "title": estimate_spec["title"],
                "json_path": "charts/estimate-results.json",
                "svg_path": "charts/estimate-results.svg",
                "png_path": "charts/estimate-results.png",
            }
        )

    workflow_source = _workflow_mermaid(project)
    (diagrams_dir / "research-workflow.mmd").write_text(
        workflow_source, encoding="utf-8"
    )
    (diagrams_dir / "research-workflow.svg").write_text(
        _workflow_svg(stages), encoding="utf-8"
    )
    _write_workflow_png(diagrams_dir / "research-workflow.png", len(stages))

    diagram_records = [
        {
            "diagram_id": "research-workflow",
            "title": "十阶段研究工作流",
            "mmd_path": "diagrams/research-workflow.mmd",
            "svg_path": "diagrams/research-workflow.svg",
            "png_path": "diagrams/research-workflow.png",
        }
    ]

    graph = _claim_graph(project)
    if graph["nodes"]:
        claim_source = _claim_mermaid(graph)
        (diagrams_dir / "claim-evidence.mmd").write_text(
            claim_source, encoding="utf-8"
        )
        (diagrams_dir / "claim-evidence.svg").write_text(
            _claim_graph_svg(graph), encoding="utf-8"
        )
        _write_claim_graph_png(
            diagrams_dir / "claim-evidence.png",
            len(graph["evidence"]),
            len(graph["claims"]),
            len(graph["conclusions"]),
        )
        diagram_records.append(
            {
                "diagram_id": "claim-evidence",
                "title": "Claim-Evidence 可追溯关系",
                "mmd_path": "diagrams/claim-evidence.mmd",
                "svg_path": "diagrams/claim-evidence.svg",
                "png_path": "diagrams/claim-evidence.png",
            }
        )

    return {
        "charts": chart_records,
        "diagrams": diagram_records,
        "estimate_items": estimates,
    }


def enhance_interactive_html(
    html_path: Path,
    export_dir: Path,
    visual_assets: dict[str, Any],
) -> None:
    chart_cards: list[str] = []
    specifications = sorted(
        {
            str(item.get("specification_id") or "main")
            for item in visual_assets.get("estimate_items", [])
        }
    )
    for chart in visual_assets.get("charts", []):
        svg = (export_dir / chart["svg_path"]).read_text(encoding="utf-8")
        controls = ""
        if chart["kind"] == "forest":
            options = "".join(
                f'<option value="{escape(value, quote=True)}">{escape(value)}</option>'
                for value in specifications
            )
            controls = f"""
              <div class="visual-controls">
                <label>规格<select data-estimate-filter><option value="all">全部</option>{options}</select></label>
                <label class="visual-toggle"><input type="checkbox" data-ci-toggle checked> 显示置信区间</label>
              </div>
            """
        chart_cards.append(
            f"""
            <article class="visual-panel" data-chart-panel="{escape(chart['chart_id'], quote=True)}">
              <header><div><span>CHART</span><h3>{escape(chart['title'])}</h3></div>{controls}</header>
              <div class="embedded-svg">{svg}</div>
            </article>
            """
        )

    diagram_cards: list[str] = []
    for diagram in visual_assets.get("diagrams", []):
        svg = (export_dir / diagram["svg_path"]).read_text(encoding="utf-8")
        source = (export_dir / diagram["mmd_path"]).read_text(encoding="utf-8")
        diagram_cards.append(
            f"""
            <article class="visual-panel diagram-panel" data-diagram-panel="{escape(diagram['diagram_id'], quote=True)}">
              <header>
                <div><span>MERMAID</span><h3>{escape(diagram['title'])}</h3></div>
                <button type="button" data-copy-mermaid>复制源码</button>
              </header>
              <div class="embedded-svg">{svg}</div>
              <details>
                <summary>查看 Mermaid 源码</summary>
                <pre data-mermaid-source>{escape(source)}</pre>
              </details>
            </article>
            """
        )

    visuals_html = f"""
      <section class="interactive-report" id="interactive-visuals" data-ai4ms-interactive-report>
        <div class="interactive-report-heading">
          <div><span>INTERACTIVE EVIDENCE</span><h2>图表与研究关系图</h2></div>
          <button type="button" onclick="window.print()">打印或另存为 PDF</button>
        </div>
        <div class="visual-stack">{''.join(chart_cards)}{''.join(diagram_cards)}</div>
      </section>
    """
    css = """
    <style>
      .interactive-report{margin-top:32px}.interactive-report-heading{display:flex;align-items:flex-end;justify-content:space-between;gap:20px;margin-bottom:18px}
      .interactive-report-heading span,.visual-panel header span{font-size:11px;font-weight:800;color:#16856b}
      .interactive-report-heading h2,.visual-panel h3{margin:4px 0 0}.interactive-report button{border:1px solid #bcc5ca;background:#fff;color:#20272b;padding:8px 12px;cursor:pointer}
      .visual-stack{display:grid;gap:18px}.visual-panel{border:1px solid #d7dde1;background:#fff;padding:16px;overflow:hidden}
      .visual-panel header{display:flex;align-items:center;justify-content:space-between;gap:16px;margin-bottom:12px}.visual-panel header>div:first-child{min-width:0}
      .visual-controls{display:flex;align-items:center;gap:12px;font-size:12px}.visual-controls label{display:flex;align-items:center;gap:6px}.visual-controls select{border:1px solid #bbc4c9;padding:6px;background:#fff}
      .visual-toggle input{accent-color:#16856b}.embedded-svg{overflow-x:auto}.embedded-svg svg{display:block;min-width:720px;width:100%;height:auto}
      .diagram-panel details{margin-top:12px;border-top:1px solid #e3e7e9;padding-top:10px}.diagram-panel summary{cursor:pointer;font-weight:700}
      .diagram-panel pre{white-space:pre-wrap;overflow-wrap:anywhere;background:#111719;color:#e6f3ee;padding:14px;font-size:12px;line-height:1.6}
      .graph-node{cursor:pointer;transition:opacity .15s}.graph-edge{transition:opacity .15s}.diagram-panel.is-focused .graph-node,.diagram-panel.is-focused .graph-edge{opacity:.18}
      .diagram-panel.is-focused .graph-node.is-related,.diagram-panel.is-focused .graph-edge.is-related{opacity:1}
      @media(max-width:760px){.interactive-report-heading,.visual-panel header{align-items:flex-start;flex-direction:column}.visual-controls{align-items:flex-start;flex-direction:column}.embedded-svg svg{min-width:620px}}
      @media print{.interactive-report button,.visual-controls{display:none}.visual-panel{break-inside:avoid}.embedded-svg svg{min-width:0}}
    </style>
    """
    script = """
    <script>
    (function(){
      var filter = document.querySelector("[data-estimate-filter]");
      var toggle = document.querySelector("[data-ci-toggle]");
      function updateEstimateChart(){
        var value = filter ? filter.value : "all";
        document.querySelectorAll("[data-estimate-row]").forEach(function(row){
          row.style.display = value === "all" || row.getAttribute("data-spec") === value ? "" : "none";
        });
        document.querySelectorAll(".ci-line").forEach(function(line){
          line.style.visibility = !toggle || toggle.checked ? "visible" : "hidden";
        });
      }
      if(filter) filter.addEventListener("change", updateEstimateChart);
      if(toggle) toggle.addEventListener("change", updateEstimateChart);
      updateEstimateChart();

      document.querySelectorAll("[data-copy-mermaid]").forEach(function(button){
        button.addEventListener("click", function(){
          var source = button.closest("[data-diagram-panel]").querySelector("[data-mermaid-source]");
          navigator.clipboard.writeText(source.textContent || "").then(function(){
            button.textContent = "已复制";
            window.setTimeout(function(){ button.textContent = "复制源码"; }, 1400);
          });
        });
      });

      document.querySelectorAll("[data-diagram-panel]").forEach(function(panel){
        panel.querySelectorAll(".graph-node").forEach(function(node){
          node.addEventListener("click", function(){
            var id = node.getAttribute("data-node");
            var already = panel.classList.contains("is-focused") && node.classList.contains("is-related");
            panel.classList.toggle("is-focused", !already);
            panel.querySelectorAll(".is-related").forEach(function(item){ item.classList.remove("is-related"); });
            if(already) return;
            node.classList.add("is-related");
            panel.querySelectorAll('.graph-edge[data-from="'+id+'"],.graph-edge[data-to="'+id+'"]').forEach(function(edge){
              edge.classList.add("is-related");
              var other = edge.getAttribute("data-from") === id ? edge.getAttribute("data-to") : edge.getAttribute("data-from");
              var related = panel.querySelector('.graph-node[data-node="'+other+'"]');
              if(related) related.classList.add("is-related");
            });
          });
        });
      });
    })();
    </script>
    """
    html = html_path.read_text(encoding="utf-8")
    html = html.replace("</head>", f"{css}</head>", 1)
    marker = '<div class="footer">'
    html = html.replace(marker, f"{visuals_html}{marker}", 1)
    html = html.replace("</body>", f"{script}</body>", 1)
    html_path.write_text(html, encoding="utf-8")


def render_docx(
    project: dict[str, Any],
    output_path: Path,
    export_dir: Path,
    visual_assets: dict[str, Any],
) -> None:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
    from docx.shared import Inches, Pt

    document = Document()
    document.core_properties.title = _report_title(project)
    document.core_properties.subject = "AI4MS 管理科学研究报告"
    document.core_properties.author = "AI4MS Research Workbench"

    normal = document.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(10.5)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    for style_name in ("Title", "Heading 1", "Heading 2", "Heading 3"):
        style = document.styles[style_name]
        style.font.name = "Arial"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")

    title = document.add_heading(_report_title(project), 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle = document.add_paragraph("AI4MS 管理科学科研工作台 · 可审计研究交付")
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER

    stages = _stage_map(project)
    delivery = stages.get("delivery", {}).get("content", {})
    evidence = stages.get("evidence", {}).get("content", {})
    robustness = stages.get("robustness", {}).get("content", {})
    analysis = stages.get("analysis", {}).get("content", {})

    document.add_heading("执行摘要", level=1)
    document.add_paragraph(str(delivery.get("executive_summary") or ""))

    document.add_heading("研究流程与审批状态", level=1)
    _docx_table(
        document,
        ["阶段", "资产", "状态", "Revision", "内容哈希"],
        [
            [
                stage.get("code", ""),
                stage.get("title", ""),
                stage.get("status", ""),
                stage.get("revision", 0),
                str(stage.get("content_hash") or "")[:12],
            ]
            for stage in project.get("stages", [])
        ],
    )
    _docx_picture(
        document,
        export_dir / "charts" / "stage-revisions.png",
        "图 1 研究阶段版本状态",
        Inches(6.4),
    )

    document.add_heading("核心结论", level=1)
    _docx_table(
        document,
        ["结论 ID", "结论", "状态", "主张 ID", "证据 ID", "适用边界"],
        [
            [
                item.get("conclusion_id", ""),
                item.get("statement", ""),
                item.get("status", ""),
                ", ".join(item.get("claim_ids", [])),
                ", ".join(item.get("evidence_ids", [])),
                item.get("scope_note", ""),
            ]
            for item in delivery.get("conclusions", [])
            if isinstance(item, dict)
        ],
    )

    document.add_heading("Claim-Evidence 可追溯矩阵", level=1)
    claims = [
        item
        for item in evidence.get("claims", [])
        if isinstance(item, dict)
    ]
    _docx_table(
        document,
        ["主张 ID", "主张", "状态", "置信度", "证据", "稳健性检查"],
        [
            [
                claim.get("claim_id", ""),
                claim.get("claim_text", ""),
                claim.get("status", ""),
                claim.get("confidence", ""),
                "; ".join(
                    f"{item.get('evidence_id', '')}:{item.get('direction', '')}@{item.get('artifact_id', '')}"
                    for item in claim.get("evidence", [])
                    if isinstance(item, dict)
                ),
                ", ".join(claim.get("robustness_check_ids", [])),
            ]
            for claim in claims
        ],
    )
    claim_image = export_dir / "diagrams" / "claim-evidence.png"
    if claim_image.is_file():
        _docx_picture(
            document,
            claim_image,
            "图 2 Claim-Evidence 关系图，Mermaid 源码随交付包提供",
            Inches(6.4),
        )

    estimates = visual_assets.get("estimate_items", [])
    if estimates:
        document.add_heading("结构化估计结果", level=1)
        _docx_picture(
            document,
            export_dir / "charts" / "estimate-results.png",
            "图 3 估计值与置信区间",
            Inches(6.4),
        )
        _docx_table(
            document,
            ["结果 ID", "规格", "变量", "估计值", "标准误", "p 值", "95% CI", "N"],
            [
                [
                    item["result_id"],
                    item["specification_id"],
                    item["term"],
                    _number_text(item.get("estimate")),
                    _number_text(item.get("std_error")),
                    _number_text(item.get("p_value")),
                    f"{_number_text(item.get('ci_lower'))} 至 {_number_text(item.get('ci_upper'))}",
                    item.get("sample_size") or "",
                ]
                for item in estimates
            ],
        )

    document.add_heading("稳健性与运行记录", level=1)
    _docx_table(
        document,
        ["检查 ID", "类别", "状态", "结果摘要", "解释影响"],
        [
            [
                item.get("check_id", ""),
                item.get("category", ""),
                item.get("status", ""),
                item.get("result_summary", ""),
                item.get("implication", ""),
            ]
            for item in robustness.get("robustness_matrix", [])
            if isinstance(item, dict)
        ],
    )
    _docx_table(
        document,
        ["Run ID", "状态", "原因", "退出码", "do-file SHA-256"],
        [
            [
                run.get("run_id", ""),
                run.get("status", ""),
                run.get("reason_code", ""),
                run.get("exit_code", ""),
                str(run.get("do_file_sha256") or "")[:16],
            ]
            for run in analysis.get("runs", [])
            if isinstance(run, dict)
        ],
    )

    document.add_heading("政策与管理含义", level=1)
    for item in delivery.get("policy_implications", []):
        if isinstance(item, dict):
            document.add_paragraph(
                f"{item.get('audience', '')}：{item.get('statement', '')} "
                f"条件：{'；'.join(item.get('conditions', []))} "
                f"风险：{item.get('risk_note', '')}",
                style="List Bullet",
            )

    document.add_heading("局限、披露与复现说明", level=1)
    for item in delivery.get("limitations", []):
        document.add_paragraph(str(item), style="List Bullet")
    document.add_paragraph(str(delivery.get("disclosure") or ""))
    for item in delivery.get("reproducibility_notes", []):
        document.add_paragraph(str(item), style="List Bullet")

    document.add_heading("参考文献", level=1)
    for paper in delivery.get("references", []):
        if isinstance(paper, dict):
            authors = ", ".join(str(value) for value in paper.get("authors", []))
            document.add_paragraph(
                f"[{paper.get('paper_id', '')}] {authors}. {paper.get('title', '')}. "
                f"{paper.get('year', '')}. DOI: {paper.get('doi', '')}",
                style="List Number",
            )
    document.save(output_path)


def render_pdf(
    project: dict[str, Any],
    output_path: Path,
    visual_assets: dict[str, Any],
) -> None:
    from PIL import Image, ImageDraw, ImageFont

    page_width, page_height = 1240, 1754
    margin = 92
    font_path = _report_font_path()
    fonts = {
        size: ImageFont.truetype(str(font_path), size=size)
        for size in (22, 26, 32, 48)
    }
    pages: list[Any] = []
    image: Any
    draw: Any
    y: int
    started = False

    def new_page() -> None:
        nonlocal image, draw, y, started
        if started:
            pages.append(image)
        image = Image.new("RGB", (page_width, page_height), "white")
        draw = ImageDraw.Draw(image)
        draw.line(
            (margin, 66, page_width - margin, 66),
            fill=(22, 133, 107),
            width=5,
        )
        y = 96
        started = True

    def text(
        value: Any,
        *,
        size: int = 22,
        indent: int = 0,
        gap: int = 12,
        color: tuple[int, int, int] = (35, 43, 47),
    ) -> None:
        nonlocal y
        font = fonts[size]
        max_width = page_width - margin * 2 - indent
        for line in _wrap_pixels(draw, str(value or ""), font, max_width):
            line_height = int(font.size * 1.45)
            if y + line_height > page_height - 88:
                new_page()
            draw.text((margin + indent, y), line, font=font, fill=color)
            y += line_height + gap

    def heading(value: str) -> None:
        nonlocal y
        if y + 80 > page_height - 88:
            new_page()
        y += 20
        text(value, size=32, gap=16, color=(20, 76, 63))

    new_page()

    stages = _stage_map(project)
    delivery = stages.get("delivery", {}).get("content", {})
    evidence = stages.get("evidence", {}).get("content", {})
    robustness = stages.get("robustness", {}).get("content", {})
    analysis = stages.get("analysis", {}).get("content", {})

    text(_report_title(project), size=48, gap=26, color=(24, 32, 36))
    text("AI4MS 管理科学科研工作台 · 固定版研究报告", size=26, gap=28, color=(86, 101, 108))
    heading("执行摘要")
    text(delivery.get("executive_summary", ""))
    heading("研究流程与审批状态")
    for stage in project.get("stages", []):
        text(
            f"{stage.get('code', '')} {stage.get('title', '')} | "
            f"{stage.get('status', '')} | Revision {stage.get('revision', 0)}",
            indent=8,
        )

    heading("核心结论")
    for item in delivery.get("conclusions", []):
        if isinstance(item, dict):
            text(
                f"{item.get('conclusion_id', '')} [{item.get('status', '')}] "
                f"{item.get('statement', '')}",
                indent=8,
            )
            text(
                f"主张：{', '.join(item.get('claim_ids', []))}；"
                f"证据：{', '.join(item.get('evidence_ids', []))}；"
                f"边界：{item.get('scope_note', '')}",
                size=22,
                indent=16,
            )

    heading("Claim-Evidence 可追溯矩阵")
    for claim in evidence.get("claims", []):
        if isinstance(claim, dict):
            text(
                f"{claim.get('claim_id', '')} [{claim.get('status', '')}/"
                f"{claim.get('confidence', '')}] {claim.get('claim_text', '')}",
                indent=8,
            )
            links = [
                f"{item.get('evidence_id', '')}:{item.get('direction', '')}@{item.get('artifact_id', '')}"
                for item in claim.get("evidence", [])
                if isinstance(item, dict)
            ]
            text("证据：" + "；".join(links), size=22, indent=16)

    estimates = visual_assets.get("estimate_items", [])
    if estimates:
        heading("结构化估计结果")
        for item in estimates:
            text(
                f"{item['result_id']} | {item['specification_id']} | {item['term']} | "
                f"estimate={_number_text(item.get('estimate'))} | "
                f"SE={_number_text(item.get('std_error'))} | "
                f"p={_number_text(item.get('p_value'))} | "
                f"95% CI [{_number_text(item.get('ci_lower'))}, {_number_text(item.get('ci_upper'))}]",
                size=22,
                indent=8,
            )

    heading("稳健性与运行记录")
    for item in robustness.get("robustness_matrix", []):
        if isinstance(item, dict):
            text(
                f"{item.get('check_id', '')} [{item.get('status', '')}] "
                f"{item.get('result_summary', '')}；{item.get('implication', '')}",
                indent=8,
            )
    for run in analysis.get("runs", []):
        if isinstance(run, dict):
            text(
                f"{run.get('run_id', '')} | {run.get('status', '')} | "
                f"{run.get('reason_code', '')} | do-file {str(run.get('do_file_sha256') or '')[:16]}",
                size=22,
                indent=8,
            )

    heading("政策与管理含义")
    for item in delivery.get("policy_implications", []):
        if isinstance(item, dict):
            text(
                f"{item.get('audience', '')}：{item.get('statement', '')} "
                f"风险：{item.get('risk_note', '')}",
                indent=8,
            )

    heading("局限、披露与复现说明")
    for item in delivery.get("limitations", []):
        text(f"• {item}", indent=8)
    text(delivery.get("disclosure", ""), indent=8)
    for item in delivery.get("reproducibility_notes", []):
        text(f"• {item}", indent=8)

    heading("参考文献")
    for paper in delivery.get("references", []):
        if isinstance(paper, dict):
            authors = ", ".join(str(value) for value in paper.get("authors", []))
            text(
                f"[{paper.get('paper_id', '')}] {authors}. {paper.get('title', '')}. "
                f"{paper.get('year', '')}. DOI: {paper.get('doi', '')}",
                size=22,
                indent=8,
            )
    pages.append(image)
    pages.append(_workflow_pdf_image(project, page_width, page_height, font_path))
    if estimates:
        pages.append(
            _estimate_pdf_image(
                estimates,
                page_width,
                page_height,
                font_path,
            )
        )
    pages[0].save(
        output_path,
        "PDF",
        resolution=120,
        save_all=True,
        append_images=pages[1:],
        quality=90,
    )


def reproduction_readme(project: dict[str, Any]) -> str:
    stages = _stage_map(project)
    analysis = stages.get("analysis", {}).get("content", {})
    identification = stages.get("identification", {}).get("content", {})
    runs = [
        run for run in analysis.get("runs", []) if isinstance(run, dict)
    ]
    lines = [
        "# AI4MS Stata 复现包",
        "",
        "本包用于复核已批准 do-file、运行契约、结构化结果、日志、图表和文件签名。",
        "",
        "## 运行边界",
        "",
        "- Stata 不包含在 Docker 镜像或本复现包内，需研究者或评委提供合法授权。",
        "- 原始 `.dta` 默认不进入交付包；请按 `data-assets.json` 中的 SHA-256 提供同一输入文件。",
        "- 推荐通过 AI4MS Local Runner 执行，以校验 input hash、do-file hash 和 Result Bundle。",
        "",
        "## 固定分析计划",
        "",
        f"- 执行引擎：{identification.get('execution_engine', 'unknown')}",
        f"- 代码语言：{identification.get('code_language', '')}",
        f"- 随机种子：{identification.get('seed', '')}",
        f"- 已批准分析计划 revision：{analysis.get('approved_analysis_plan_revision', 0)}",
        f"- 已批准分析计划 hash：{analysis.get('approved_analysis_plan_hash', '')}",
        "",
        "## 复现步骤",
        "",
        "1. 安装并授权与 runner-profile.json 相符的 Stata。",
        "2. 准备与 data-assets.json 中 SHA-256 一致的 `.dta` 文件。",
        "3. 启动 `ai4ms-stata-runner`，并配置相同的 Runner Token。",
        "4. 使用包内 `analysis.do` 和运行契约提交任务。",
        "5. 对照 `result_bundle.json`、`structured_results.csv`、日志和 manifest 复核结果。",
        "",
        "## 已登记运行",
        "",
    ]
    if not runs:
        lines.append("- 当前尚无正式运行记录；包内只包含获批 do-file 与数据资产元信息。")
    for run in runs:
        lines.append(
            f"- `{run.get('run_id', '')}`：{run.get('status', '')}/"
            f"{run.get('reason_code', '')}，do-file SHA-256 "
            f"`{run.get('do_file_sha256', '')}`"
        )
    return "\n".join(lines).strip() + "\n"


def _collect_estimates(project: dict[str, Any]) -> list[dict[str, Any]]:
    analysis = _stage_map(project).get("analysis", {}).get("content", {})
    found: list[dict[str, Any]] = []
    for run in analysis.get("runs", []):
        if not isinstance(run, dict) or run.get("status") != "succeeded":
            continue
        for index, item in enumerate(run.get("structured_results", []), start=1):
            if not isinstance(item, dict):
                continue
            estimate = _as_number(item.get("estimate", item.get("value")))
            if estimate is None:
                continue
            kind = str(item.get("kind") or "estimate")
            if kind != "estimate":
                continue
            std_error = _as_number(item.get("std_error"))
            ci_lower = _as_number(item.get("ci_lower"))
            ci_upper = _as_number(item.get("ci_upper"))
            if ci_lower is None and std_error is not None:
                ci_lower = estimate - 1.96 * std_error
            if ci_upper is None and std_error is not None:
                ci_upper = estimate + 1.96 * std_error
            found.append(
                {
                    "run_id": str(run.get("run_id") or ""),
                    "result_id": str(item.get("result_id") or f"RES_{index}"),
                    "specification_id": str(item.get("specification_id") or "main"),
                    "term": str(item.get("term") or item.get("name") or item.get("label") or ""),
                    "estimate": estimate,
                    "std_error": std_error,
                    "p_value": _as_number(item.get("p_value")),
                    "ci_lower": ci_lower,
                    "ci_upper": ci_upper,
                    "sample_size": item.get("sample_size"),
                }
            )
    return found[:40]


def _stage_chart_svg(spec: dict[str, Any]) -> str:
    items = spec["items"]
    width = 1000
    row_height = 42
    height = max(240, 90 + row_height * len(items))
    max_revision = max([int(item.get("revision") or 0) for item in items] + [1])
    rows = []
    for index, item in enumerate(items):
        y = 68 + index * row_height
        bar_width = max(8, int((int(item["revision"]) / max_revision) * 560))
        color = STATUS_COLORS.get(item["status"], STATUS_COLORS["not_started"])
        rows.append(
            f'<g data-stage="{escape(item["code"], quote=True)}">'
            f'<text x="22" y="{y + 17}" class="label">{escape(item["code"])} {escape(item["title"])}</text>'
            f'<rect x="230" y="{y}" width="600" height="24" rx="3" fill="#edf0f2"/>'
            f'<rect x="230" y="{y}" width="{bar_width}" height="24" rx="3" fill="{color}"/>'
            f'<text x="846" y="{y + 17}" class="value">Revision {item["revision"]} · {escape(item["status"])}</text>'
            "</g>"
        )
    return _svg_shell(
        width,
        height,
        f'<text x="22" y="34" class="title">{escape(spec["title"])}</text>{"".join(rows)}',
    )


def _estimate_chart_svg(spec: dict[str, Any]) -> str:
    items = spec["items"]
    values = [
        value
        for item in items
        for value in (item.get("ci_lower"), item.get("ci_upper"), item.get("estimate"), 0)
        if isinstance(value, (int, float)) and math.isfinite(value)
    ]
    lower = min(values) if values else -1
    upper = max(values) if values else 1
    if math.isclose(lower, upper):
        lower -= 1
        upper += 1
    padding = (upper - lower) * 0.12
    lower -= padding
    upper += padding
    plot_left = 275
    plot_width = 650
    row_height = 44
    height = max(260, 110 + row_height * len(items))

    def x(value: float) -> float:
        return plot_left + (value - lower) / (upper - lower) * plot_width

    zero_x = x(0)
    rows = [
        f'<line x1="{zero_x:.1f}" y1="60" x2="{zero_x:.1f}" y2="{height - 32}" stroke="#8c969c" stroke-dasharray="4 4"/>'
    ]
    for index, item in enumerate(items):
        y = 82 + index * row_height
        estimate = float(item["estimate"])
        low = item.get("ci_lower")
        high = item.get("ci_upper")
        spec_id = escape(str(item["specification_id"]), quote=True)
        label = f'{item["result_id"]} · {item["term"]}'
        ci = ""
        if isinstance(low, (int, float)) and isinstance(high, (int, float)):
            ci = (
                f'<line class="ci-line" x1="{x(float(low)):.1f}" y1="{y}" '
                f'x2="{x(float(high)):.1f}" y2="{y}" stroke="#31515f" stroke-width="3"/>'
                f'<line class="ci-line" x1="{x(float(low)):.1f}" y1="{y - 5}" '
                f'x2="{x(float(low)):.1f}" y2="{y + 5}" stroke="#31515f"/>'
                f'<line class="ci-line" x1="{x(float(high)):.1f}" y1="{y - 5}" '
                f'x2="{x(float(high)):.1f}" y2="{y + 5}" stroke="#31515f"/>'
            )
        rows.append(
            f'<g data-estimate-row data-spec="{spec_id}">'
            f'<text x="18" y="{y + 5}" class="label">{escape(label[:38])}</text>'
            f'{ci}<circle cx="{x(estimate):.1f}" cy="{y}" r="6" fill="#16856b">'
            f'<title>{escape(label)}: {estimate:.5g}</title></circle>'
            f'<text x="940" y="{y + 5}" class="value">{estimate:.5g}</text>'
            "</g>"
        )
    rows.append(
        f'<text x="{plot_left}" y="{height - 10}" class="axis">{lower:.4g}</text>'
        f'<text x="{plot_left + plot_width - 24}" y="{height - 10}" class="axis">{upper:.4g}</text>'
    )
    return _svg_shell(
        1000,
        height,
        f'<text x="18" y="34" class="title">{escape(spec["title"])}</text>{"".join(rows)}',
    )


def _workflow_mermaid(project: dict[str, Any]) -> str:
    stages = [
        stage for stage in project.get("stages", []) if isinstance(stage, dict)
    ]
    lines = ["flowchart LR"]
    for stage in stages:
        code = _mermaid_id(str(stage.get("code") or "stage"))
        label = _mermaid_label(f"{stage.get('code', '')} {stage.get('title', '')}")
        lines.append(f'  {code}["{label}"]')
    for left, right in zip(stages, stages[1:]):
        lines.append(
            f"  {_mermaid_id(str(left.get('code') or 'left'))} --> "
            f"{_mermaid_id(str(right.get('code') or 'right'))}"
        )
    lines.extend(
        [
            "  classDef approved fill:#d9f3e8,stroke:#16856b,color:#102c23",
            "  classDef current fill:#fff0d5,stroke:#d48a16,color:#49310a",
            "  classDef pending fill:#eef1f3,stroke:#8a969e,color:#283238",
        ]
    )
    for stage in stages:
        status = str(stage.get("status") or "not_started")
        class_name = "approved" if status == "approved" else "current" if status in {"needs_review", "in_progress"} else "pending"
        lines.append(f"  class {_mermaid_id(str(stage.get('code') or 'stage'))} {class_name}")
    return "\n".join(lines) + "\n"


def _claim_graph(project: dict[str, Any]) -> dict[str, Any]:
    stages = _stage_map(project)
    evidence_content = stages.get("evidence", {}).get("content", {})
    delivery = stages.get("delivery", {}).get("content", {})
    evidence_nodes: dict[str, dict[str, str]] = {}
    claims: list[dict[str, Any]] = []
    edges: list[dict[str, str]] = []
    for claim in evidence_content.get("claims", []):
        if not isinstance(claim, dict) or not claim.get("claim_id"):
            continue
        claim_id = str(claim["claim_id"])
        claims.append(
            {
                "id": claim_id,
                "label": str(claim.get("claim_text") or claim_id),
                "status": str(claim.get("status") or ""),
            }
        )
        for item in claim.get("evidence", []):
            if not isinstance(item, dict) or not item.get("evidence_id"):
                continue
            evidence_id = str(item["evidence_id"])
            evidence_nodes.setdefault(
                evidence_id,
                {
                    "id": evidence_id,
                    "label": str(item.get("artifact_id") or item.get("locator") or evidence_id),
                    "status": str(item.get("direction") or ""),
                },
            )
            edges.append({"from": evidence_id, "to": claim_id})
    conclusions = []
    for item in delivery.get("conclusions", []):
        if not isinstance(item, dict) or not item.get("conclusion_id"):
            continue
        conclusion_id = str(item["conclusion_id"])
        conclusions.append(
            {
                "id": conclusion_id,
                "label": str(item.get("statement") or conclusion_id),
                "status": str(item.get("status") or ""),
            }
        )
        for claim_id in item.get("claim_ids", []):
            edges.append({"from": str(claim_id), "to": conclusion_id})
    evidence = list(evidence_nodes.values())
    return {
        "evidence": evidence,
        "claims": claims,
        "conclusions": conclusions,
        "nodes": [*evidence, *claims, *conclusions],
        "edges": edges,
    }


def _claim_mermaid(graph: dict[str, Any]) -> str:
    lines = ["flowchart LR", "  subgraph Evidence", "    direction TB"]
    for item in graph["evidence"]:
        lines.append(
            f'    E_{_mermaid_id(item["id"])}["{_mermaid_label(item["id"] + " " + item["label"])}"]'
        )
    lines.extend(["  end", "  subgraph Claims", "    direction TB"])
    for item in graph["claims"]:
        lines.append(
            f'    C_{_mermaid_id(item["id"])}["{_mermaid_label(item["id"] + " " + item["label"])}"]'
        )
    lines.extend(["  end", "  subgraph Conclusions", "    direction TB"])
    for item in graph["conclusions"]:
        lines.append(
            f'    O_{_mermaid_id(item["id"])}["{_mermaid_label(item["id"] + " " + item["label"])}"]'
        )
    lines.append("  end")
    evidence_ids = {item["id"] for item in graph["evidence"]}
    claim_ids = {item["id"] for item in graph["claims"]}
    for edge in graph["edges"]:
        prefix_from = "E_" if edge["from"] in evidence_ids else "C_"
        prefix_to = "C_" if edge["to"] in claim_ids else "O_"
        lines.append(
            f"  {prefix_from}{_mermaid_id(edge['from'])} --> "
            f"{prefix_to}{_mermaid_id(edge['to'])}"
        )
    return "\n".join(lines) + "\n"


def _workflow_svg(stages: list[dict[str, Any]]) -> str:
    width = 1200
    box_width = 96
    gap = 20
    y = 90
    elements = [
        '<defs><marker id="arrow" markerWidth="8" markerHeight="8" refX="7" refY="4" orient="auto"><path d="M0,0 L8,4 L0,8 Z" fill="#68757c"/></marker></defs>',
        '<text x="24" y="34" class="title">十阶段研究工作流</text>',
    ]
    for index, stage in enumerate(stages):
        x = 24 + index * (box_width + gap)
        color = STATUS_COLORS.get(stage["status"], STATUS_COLORS["not_started"])
        if index < len(stages) - 1:
            elements.append(
                f'<line x1="{x + box_width}" y1="{y + 31}" x2="{x + box_width + gap - 3}" y2="{y + 31}" stroke="#68757c" marker-end="url(#arrow)"/>'
            )
        elements.append(
            f'<g class="graph-node" data-node="{escape(stage["code"], quote=True)}">'
            f'<rect x="{x}" y="{y}" width="{box_width}" height="62" rx="4" fill="{color}" opacity=".16" stroke="{color}" stroke-width="2"/>'
            f'<text x="{x + box_width / 2}" y="{y + 27}" text-anchor="middle" class="node-title">{escape(stage["code"])}</text>'
            f'<text x="{x + box_width / 2}" y="{y + 47}" text-anchor="middle" class="node-label">{escape(stage["title"][:8])}</text>'
            "</g>"
        )
    return _svg_shell(width, 230, "".join(elements))


def _claim_graph_svg(graph: dict[str, Any]) -> str:
    columns = [
        ("Evidence", graph["evidence"], 40, "#2878b5"),
        ("Claims", graph["claims"], 390, "#16856b"),
        ("Conclusions", graph["conclusions"], 740, "#d48a16"),
    ]
    max_count = max([len(items) for _name, items, _x, _color in columns] + [1])
    height = max(300, 110 + max_count * 90)
    positions: dict[str, tuple[int, int]] = {}
    elements = [
        '<defs><marker id="graph-arrow" markerWidth="8" markerHeight="8" refX="7" refY="4" orient="auto"><path d="M0,0 L8,4 L0,8 Z" fill="#77858c"/></marker></defs>',
        '<text x="24" y="34" class="title">Claim-Evidence 可追溯关系</text>',
    ]
    for name, items, x, color in columns:
        elements.append(f'<text x="{x}" y="70" class="column-title">{name}</text>')
        for index, item in enumerate(items):
            y = 88 + index * 90
            positions[item["id"]] = (x, y)
            label = f'{item["id"]} · {item["label"]}'
            elements.append(
                f'<g class="graph-node" data-node="{escape(item["id"], quote=True)}">'
                f'<rect x="{x}" y="{y}" width="260" height="60" rx="4" fill="{color}" opacity=".12" stroke="{color}" stroke-width="2"/>'
                f'<text x="{x + 12}" y="{y + 24}" class="node-title">{escape(item["id"])}</text>'
                f'<text x="{x + 12}" y="{y + 44}" class="node-label">{escape(label[:34])}</text>'
                "</g>"
            )
    for edge in graph["edges"]:
        start = positions.get(edge["from"])
        end = positions.get(edge["to"])
        if not start or not end:
            continue
        elements.insert(
            2,
            f'<path class="graph-edge" data-from="{escape(edge["from"], quote=True)}" data-to="{escape(edge["to"], quote=True)}" '
            f'd="M {start[0] + 260} {start[1] + 30} C {start[0] + 310} {start[1] + 30}, '
            f'{end[0] - 50} {end[1] + 30}, {end[0] - 4} {end[1] + 30}" '
            f'fill="none" stroke="#77858c" stroke-width="1.5" marker-end="url(#graph-arrow)"/>',
        )
    return _svg_shell(1040, height, "".join(elements))


def _svg_shell(width: int, height: int, body: str) -> str:
    return (
        f'<svg xmlns="http://www.w3.org/2000/svg" width="{width}" height="{height}" '
        f'viewBox="0 0 {width} {height}" role="img">'
        "<style>"
        ".title{font:700 20px Arial,'Microsoft YaHei',sans-serif;fill:#1b252a}"
        ".label{font:13px Arial,'Microsoft YaHei',sans-serif;fill:#29343a}"
        ".value,.axis{font:12px Arial,sans-serif;fill:#59666d}"
        ".node-title{font:700 13px Arial,'Microsoft YaHei',sans-serif;fill:#172126}"
        ".node-label{font:11px Arial,'Microsoft YaHei',sans-serif;fill:#354249}"
        ".column-title{font:700 13px Arial,sans-serif;fill:#65727a}"
        "</style>"
        f'<rect width="{width}" height="{height}" fill="#ffffff"/>{body}</svg>'
    )


def _write_progress_png(path: Path, stages: list[dict[str, Any]]) -> None:
    from PIL import Image, ImageDraw, ImageFont

    image = Image.new("RGB", (1000, 500), "white")
    draw = ImageDraw.Draw(image)
    font = ImageFont.truetype(str(_report_font_path()), 18)
    maximum = max([int(item.get("revision") or 0) for item in stages] + [1])
    for index, item in enumerate(stages[:10]):
        y = 34 + index * 43
        draw.text(
            (18, y + 1),
            f"{item.get('code', '')} {str(item.get('title', ''))[:8]}",
            font=font,
            fill=(35, 43, 47),
        )
        draw.rectangle((220, y, 820, y + 24), fill=(232, 237, 239))
        width = max(8, int(600 * int(item.get("revision") or 0) / maximum))
        draw.rectangle(
            (220, y, 220 + width, y + 24),
            fill=_hex_rgb(STATUS_COLORS.get(item["status"], "#aeb6bd")),
        )
        draw.text(
            (836, y + 1),
            f"R{item.get('revision', 0)}",
            font=font,
            fill=(78, 90, 96),
        )
    image.save(path, "PNG", optimize=True)


def _write_estimate_png(path: Path, items: list[dict[str, Any]]) -> None:
    from PIL import Image, ImageDraw, ImageFont

    height = max(320, 90 + min(len(items), 24) * 34)
    image = Image.new("RGB", (1000, height), "white")
    draw = ImageDraw.Draw(image)
    font = ImageFont.truetype(str(_report_font_path()), 16)
    values = [
        value
        for item in items
        for value in (item.get("ci_lower"), item.get("ci_upper"), item.get("estimate"), 0)
        if isinstance(value, (int, float)) and math.isfinite(value)
    ]
    lower = min(values) if values else -1
    upper = max(values) if values else 1
    if math.isclose(lower, upper):
        lower -= 1
        upper += 1
    span = upper - lower
    lower -= span * 0.1
    upper += span * 0.1

    def x(value: float) -> int:
        return int(260 + (value - lower) / (upper - lower) * 610)

    draw.line((x(0), 30, x(0), height - 30), fill=(145, 154, 160), width=1)
    for index, item in enumerate(items[:24]):
        y = 55 + index * 34
        draw.text(
            (16, y - 10),
            f"{item['result_id']} {item['term']}"[:28],
            font=font,
            fill=(35, 43, 47),
        )
        low = item.get("ci_lower")
        high = item.get("ci_upper")
        if isinstance(low, (int, float)) and isinstance(high, (int, float)):
            draw.line((x(float(low)), y, x(float(high)), y), fill=(49, 81, 95), width=3)
            draw.line((x(float(low)), y - 5, x(float(low)), y + 5), fill=(49, 81, 95))
            draw.line((x(float(high)), y - 5, x(float(high)), y + 5), fill=(49, 81, 95))
        px = x(float(item["estimate"]))
        draw.rectangle((px - 5, y - 5, px + 5, y + 5), fill=(22, 133, 107))
        draw.text(
            (888, y - 10),
            _number_text(item["estimate"]),
            font=font,
            fill=(78, 90, 96),
        )
    image.save(path, "PNG", optimize=True)


def _write_workflow_png(path: Path, count: int) -> None:
    from PIL import Image, ImageDraw, ImageFont

    image = Image.new("RGB", (1200, 220), "white")
    draw = ImageDraw.Draw(image)
    font = ImageFont.truetype(str(_report_font_path()), 17)
    count = max(1, min(count, 10))
    box_width = 92
    gap = 22
    for index in range(count):
        x = 28 + index * (box_width + gap)
        if index:
            draw.line((x - gap, 110, x, 110), fill=(98, 111, 118), width=2)
        draw.rectangle(
            (x, 78, x + box_width, 142),
            fill=(217, 243, 232),
            outline=(22, 133, 107),
            width=2,
        )
        draw.text((x + 28, 98), f"S{index}", font=font, fill=(27, 46, 40))
    image.save(path, "PNG", optimize=True)


def _write_claim_graph_png(path: Path, evidence: int, claims: int, conclusions: int) -> None:
    from PIL import Image, ImageDraw, ImageFont

    height = max(340, 80 + max(evidence, claims, conclusions, 1) * 72)
    image = Image.new("RGB", (1000, height), "white")
    draw = ImageDraw.Draw(image)
    font = ImageFont.truetype(str(_report_font_path()), 17)
    columns = [
        ("Evidence", 40, evidence, (226, 239, 248)),
        ("Claims", 370, claims, (217, 243, 232)),
        ("Conclusions", 700, conclusions, (255, 239, 213)),
    ]
    positions: list[list[tuple[int, int]]] = []
    for title, x, count, color in columns:
        draw.text((x, 12), title, font=font, fill=(78, 90, 96))
        column = []
        for index in range(max(1, count)):
            y = 44 + index * 72
            draw.rectangle(
                (x, y, x + 250, y + 48),
                fill=color,
                outline=(92, 106, 114),
                width=2,
            )
            prefix = "EV" if title == "Evidence" else "C" if title == "Claims" else "CON"
            draw.text(
                (x + 12, y + 12),
                f"{prefix}{index + 1}",
                font=font,
                fill=(35, 43, 47),
            )
            column.append((x, y + 24))
        positions.append(column)
    for left, right in zip(positions, positions[1:]):
        for index, start in enumerate(left):
            end = right[min(index, len(right) - 1)]
            draw.line(
                (start[0] + 250, start[1], end[0], end[1]),
                fill=(116, 130, 137),
                width=2,
            )
    image.save(path, "PNG", optimize=True)


def _report_font_path() -> Path:
    configured = os.environ.get("AI4MS_REPORT_FONT", "").strip()
    candidates = [
        configured,
        "C:/Windows/Fonts/msyh.ttc",
        "C:/Windows/Fonts/simhei.ttf",
        "/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc",
        "/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.otf",
        "/usr/share/fonts/truetype/noto/NotoSansCJK-Regular.ttc",
    ]
    for value in candidates:
        if value and Path(value).is_file():
            return Path(value)
    raise RuntimeError(
        "未找到可用于中文报告的字体；请安装 fonts-noto-cjk 或设置 AI4MS_REPORT_FONT"
    )


def _wrap_pixels(draw: Any, value: str, font: Any, width: int) -> list[str]:
    text = str(value or "").replace("\r", " ").replace("\n", " ").strip()
    if not text:
        return [""]
    lines: list[str] = []
    current = ""
    for char in text:
        candidate = current + char
        if current and draw.textlength(candidate, font=font) > width:
            lines.append(current)
            current = char
        else:
            current = candidate
    if current:
        lines.append(current)
    return lines


def _workflow_pdf_image(
    project: dict[str, Any],
    width: int,
    height: int,
    font_path: Path,
) -> Any:
    from PIL import Image, ImageDraw, ImageFont

    image = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(image)
    title_font = ImageFont.truetype(str(font_path), 42)
    node_font = ImageFont.truetype(str(font_path), 23)
    draw.text((86, 86), "十阶段研究工作流", font=title_font, fill=(25, 34, 38))
    stages = [
        stage for stage in project.get("stages", []) if isinstance(stage, dict)
    ][:10]
    box_width, box_height = 188, 112
    for index, stage in enumerate(stages):
        row, column = divmod(index, 5)
        x = 86 + column * 220
        y = 260 + row * 300
        color = _hex_rgb(
            STATUS_COLORS.get(str(stage.get("status")), "#aeb6bd")
        )
        draw.rounded_rectangle(
            (x, y, x + box_width, y + box_height),
            radius=8,
            fill=tuple(int(channel + (255 - channel) * 0.82) for channel in color),
            outline=color,
            width=4,
        )
        draw.text(
            (x + 18, y + 20),
            str(stage.get("code") or ""),
            font=node_font,
            fill=(25, 38, 35),
        )
        draw.text(
            (x + 18, y + 62),
            str(stage.get("title") or "")[:8],
            font=node_font,
            fill=(43, 57, 53),
        )
        if column < 4 and index < len(stages) - 1:
            draw.line(
                (x + box_width, y + 56, x + 218, y + 56),
                fill=(105, 118, 124),
                width=3,
            )
    return image


def _estimate_pdf_image(
    items: list[dict[str, Any]],
    width: int,
    height: int,
    font_path: Path,
) -> Any:
    from PIL import Image, ImageDraw, ImageFont

    image = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(image)
    title_font = ImageFont.truetype(str(font_path), 42)
    font = ImageFont.truetype(str(font_path), 20)
    draw.text(
        (86, 86),
        "结构化估计结果与置信区间",
        font=title_font,
        fill=(25, 34, 38),
    )
    values = [
        value
        for item in items
        for value in (
            item.get("ci_lower"),
            item.get("ci_upper"),
            item.get("estimate"),
            0,
        )
        if isinstance(value, (int, float))
    ]
    lower, upper = min(values), max(values)
    if math.isclose(lower, upper):
        lower -= 1
        upper += 1
    span = upper - lower
    lower -= span * 0.1
    upper += span * 0.1

    def x(value: float) -> int:
        return int(380 + (value - lower) / (upper - lower) * 680)

    draw.line((x(0), 200, x(0), height - 120), fill=(145, 154, 160), width=2)
    for index, item in enumerate(items[:28]):
        y = 230 + index * 48
        draw.text(
            (86, y - 13),
            f"{item['result_id']} {item['term']}"[:28],
            font=font,
            fill=(35, 43, 47),
        )
        low, high = item.get("ci_lower"), item.get("ci_upper")
        if isinstance(low, (int, float)) and isinstance(high, (int, float)):
            draw.line(
                (x(float(low)), y, x(float(high)), y),
                fill=(49, 81, 95),
                width=4,
            )
        px = x(float(item["estimate"]))
        draw.rectangle((px - 7, y - 7, px + 7, y + 7), fill=(22, 133, 107))
        draw.text(
            (1080, y - 13),
            _number_text(item["estimate"]),
            font=font,
            fill=(78, 90, 96),
        )
    return image


def _docx_table(document: Any, headers: list[str], rows: list[list[Any]]) -> None:
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for cell, header in zip(table.rows[0].cells, headers):
        cell.text = str(header)
    if not rows:
        row = table.add_row().cells
        row[0].text = "暂无记录"
        if len(row) > 1:
            row[0].merge(row[-1])
    for values in rows:
        cells = table.add_row().cells
        for cell, value in zip(cells, values):
            cell.text = str(value if value is not None else "")


def _docx_picture(document: Any, path: Path, caption: str, width: Any) -> None:
    if not path.is_file():
        return
    paragraph = document.add_paragraph()
    paragraph.alignment = 1
    paragraph.add_run().add_picture(str(path), width=width)
    caption_paragraph = document.add_paragraph(caption)
    caption_paragraph.alignment = 1


def _stage_map(project: dict[str, Any]) -> dict[str, dict[str, Any]]:
    return {
        str(stage.get("key")): stage
        for stage in project.get("stages", [])
        if isinstance(stage, dict) and stage.get("key")
    }


def _report_title(project: dict[str, Any]) -> str:
    delivery = _stage_map(project).get("delivery", {}).get("content", {})
    return str(delivery.get("title") or project.get("title") or "AI4MS 研究报告")


def _write_json(path: Path, value: Any) -> None:
    path.write_text(
        json.dumps(value, ensure_ascii=False, indent=2, sort_keys=True),
        encoding="utf-8",
    )


def _as_number(value: Any) -> float | None:
    if value is None or value == "":
        return None
    try:
        number = float(value)
    except (TypeError, ValueError):
        return None
    return number if math.isfinite(number) else None


def _number_text(value: Any) -> str:
    number = _as_number(value)
    return "—" if number is None else f"{number:.6g}"


def _mermaid_id(value: str) -> str:
    return re.sub(r"[^0-9A-Za-z_]", "_", value) or "node"


def _mermaid_label(value: str) -> str:
    return (
        str(value)
        .replace('"', "'")
        .replace("\r", " ")
        .replace("\n", " ")
        .replace("[", "(")
        .replace("]", ")")[:120]
    )


def _hex_rgb(value: str) -> tuple[int, int, int]:
    value = value.lstrip("#")
    return tuple(int(value[index : index + 2], 16) for index in (0, 2, 4))  # type: ignore[return-value]
